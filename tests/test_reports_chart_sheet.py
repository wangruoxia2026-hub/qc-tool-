"""
月报质控图 sheet 测试（用户需求 2026-10-12 第二次补充）：
build_monthly_report_xlsx 每指标组新增第 4 个 sheet「质控图」——
L+H 成对时嵌入一图双 y 轴（左轴 L、右轴 H，控制限重合）；单水平时嵌入单图。

Run:
    python -m pytest tests/test_reports_chart_sheet.py -v
"""
from __future__ import annotations

import io
import sqlite3
from pathlib import Path

import pytest
from openpyxl import load_workbook

from qc_tool.db import init_database, get_connection
from qc_tool import projects as P
from qc_tool import reports as R


# ============================================================================
# Fixtures / helpers
# ============================================================================

@pytest.fixture
def conn(tmp_path: Path) -> sqlite3.Connection:
    db = tmp_path / "qc.db"
    init_database(db)
    c = get_connection(db)
    c.execute("INSERT INTO lab_settings(id, lab_name) VALUES (1, 'Test Lab')")
    c.commit()
    try:
        yield c
    finally:
        c.close()


def _seed_pair(conn, *, project="生化全套", name="ALT", lot="LOT2026A"):
    p = P.create_project(
        conn, name=project, short="SH", unit="U/L", manufacturer="Roche")
    ti_l = P.create_test_item(
        conn, p.id, name=name, lot=lot, level="L",
        target=10.0, sd=1.0, short="ALT-S", unit="U/L",
        manufacturer="Roche", operator="Op-1", pre_processor="Pre-1")
    ti_h = P.create_test_item(
        conn, p.id, name=name, lot=lot, level="H",
        target=50.0, sd=5.0, short="ALT-S", unit="U/L",
        manufacturer="Roche", operator="Op-1", pre_processor="Pre-1")
    conn.commit()
    return p, ti_l, ti_h


def _insert(conn, ti_id, pid, date, value, level, lot="LOT2026A"):
    conn.execute(
        "INSERT INTO measurements("
        "  test_item_id, project_id, lot, level, date, seq, value,"
        "  status, rules"
        ") VALUES (?, ?, ?, ?, ?, 0, ?, 'InControl', NULL)",
        (ti_id, pid, lot, level, date, value),
    )
    conn.commit()


# ============================================================================
# Tests
# ============================================================================

class TestMonthlyChartSheet:
    def test_pair_chart_sheet_with_image(self, conn):
        """L+H 成对 → 「质控图」sheet 存在且嵌入图片。"""
        p, ti_l, ti_h = _seed_pair(conn)
        _insert(conn, ti_l.id, p.id, "2026-05-01", 10.5, "L")
        _insert(conn, ti_l.id, p.id, "2026-05-02", 9.8, "L")
        _insert(conn, ti_h.id, p.id, "2026-05-01", 51.0, "H")
        _insert(conn, ti_h.id, p.id, "2026-05-02", 48.0, "H")

        xlsx = R.build_monthly_report_xlsx(
            conn, project_name=p.name, year=2026, month=5,
            selected_test_item_ids=[ti_l.id, ti_h.id],
        )
        wb = load_workbook(io.BytesIO(xlsx))
        sheet_names = wb.sheetnames

        # 每指标 4 个 sheet：本月原始 / 累计统计 / 去失控累计 / 质控图
        assert "ALT-L - 质控图" in sheet_names, sheet_names
        assert "ALT-H - 质控图" in sheet_names, sheet_names

        # 质控图 sheet 内嵌图片
        ws = wb["ALT-L - 质控图"]
        assert len(ws._images) >= 1, "质控图 sheet 应嵌入图片"
        assert ws.cell(row=1, column=1).value.startswith("生化全套")

    def test_single_level_chart_sheet(self, conn):
        """只有 L 水平 → 质控图 sheet 嵌单水平图，不报错。"""
        p = P.create_project(
            conn, name="单水平项目", short="ONE", unit="U/L", manufacturer="M")
        ti = P.create_test_item(
            conn, p.id, name="GLU", lot="LOT_S", level="L",
            target=5.0, sd=0.5, short="GLU-S", unit="mmol/L",
            manufacturer="M", operator="Op", pre_processor="Pre")
        conn.commit()
        _insert(conn, ti.id, p.id, "2026-05-01", 5.2, "L", lot="LOT_S")

        xlsx = R.build_monthly_report_xlsx(
            conn, project_name=p.name, year=2026, month=5,
            selected_test_item_ids=[ti.id],
        )
        wb = load_workbook(io.BytesIO(xlsx))
        assert "GLU-L - 质控图" in wb.sheetnames
        assert len(wb["GLU-L - 质控图"]._images) >= 1

    def test_chart_sheet_empty_month_no_crash(self, conn):
        """当月无数据 → 质控图 sheet 仍生成（嵌入 No Data 占位图）。"""
        p, ti_l, ti_h = _seed_pair(conn)

        xlsx = R.build_monthly_report_xlsx(
            conn, project_name=p.name, year=2026, month=7,
            selected_test_item_ids=[ti_l.id],
        )
        wb = load_workbook(io.BytesIO(xlsx))
        assert "ALT-L - 质控图" in wb.sheetnames
        assert len(wb["ALT-L - 质控图"]._images) >= 1
