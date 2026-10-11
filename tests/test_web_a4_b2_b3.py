"""
§A.4 / §B.2 / §B.3 后端 API 测试

覆盖：
- §A.4  POST /api/test_items/<tid>/lots 加新批号
- §B.3 GET /api/template.xlsx 下载模板
- §B.2 POST /api/test_items/<tid>/import 批量导入

跑法：
    python -m pytest tests/test_web_a4_b2_b3.py -v
"""
from __future__ import annotations

import io
import json
import socket
import threading
import time
import zipfile
from http.server import ThreadingHTTPServer
from pathlib import Path
from urllib import request as urlreq
from urllib.error import HTTPError

import pytest


def _free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


@pytest.fixture
def web_server(tmp_path: Path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    from qc_tool import web as webmod
    port = _free_port()
    handler = webmod.make_handler(db_path=tmp_path / "qc.db")
    server = ThreadingHTTPServer(("127.0.0.1", port), handler)
    server.db_path = tmp_path / "qc.db"
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base = f"http://127.0.0.1:{port}"
    time.sleep(0.05)
    try:
        yield base, tmp_path
    finally:
        server.shutdown()
        server.server_close()


def _http(method: str, url: str, data=None, raw=None, headers=None) -> tuple[int, object, dict]:
    if data is not None:
        body = json.dumps(data).encode("utf-8")
        h = {"Content-Type": "application/json"}
    elif raw is not None:
        body = raw
        h = {"Content-Type": "application/octet-stream"}
    else:
        body = b""
        h = {}
    if headers:
        h.update(headers)
    req = urlreq.Request(url, data=body, method=method, headers=h)
    try:
        resp = urlreq.urlopen(req, timeout=5)
        return resp.status, resp.read(), dict(resp.headers)
    except HTTPError as e:
        return e.code, e.read() if e.read else b"", dict(e.headers or {})


def _seed(base: str) -> dict:
    _, p, _ = _http("POST", f"{base}/api/projects", data={"name": "维生素 D"})
    p = json.loads(p) if isinstance(p, bytes) else p
    _, ti_l, _ = _http("POST", f"{base}/api/projects/{p['id']}/test_items",
                       data={"name": "D2", "lot": "LOT1", "level": "L",
                             "target": 100.0, "sd": 2.0})
    ti_l = json.loads(ti_l) if isinstance(ti_l, bytes) else ti_l
    return {"pid": p["id"], "tid_l": ti_l["id"]}


# ============================================================================
# §A.4 添加新批号
# ============================================================================

class TestA4AddLot:
    """spec §A.4：添加新批号 → 质控图区出现新批号曲线。"""

    def test_01_加新批号_成功(self, web_server):
        base, _ = web_server
        ids = _seed(base)
        status, body, _ = _http(
            "POST", f"{base}/api/test_items/{ids['tid_l']}/lots",
            data={"new_lot": "LOT2", "target": 100.0, "sd": 2.0},
        )
        assert status == 200, f"加批号应 200，实际 {status}: {body}"
        new_ti = json.loads(body) if isinstance(body, bytes) else body
        assert new_ti["lot"] == "LOT2"
        assert new_ti["level"] == "L"
        assert new_ti["id"] != ids["tid_l"]
    def test_02_加批号后列出能看到(self, web_server):
        base, _ = web_server
        ids = _seed(base)
        _http("POST", f"{base}/api/test_items/{ids['tid_l']}/lots",
              data={"new_lot": "LOT99", "target": 100.0, "sd": 2.0})
        # GET 列出 → 应能看到 LOT1 和 LOT99
        status, body, _ = _http("GET", f"{base}/api/projects/{ids['pid']}/test_items")
        items = json.loads(body)
        lots = [it["lot"] for it in items]
        assert "LOT1" in lots and "LOT99" in lots

    def test_03_加批号SD为0_400(self, web_server):
        base, _ = web_server
        ids = _seed(base)
        status, body, _ = _http(
            "POST", f"{base}/api/test_items/{ids['tid_l']}/lots",
            data={"new_lot": "LOTX", "target": 100.0, "sd": 0},
        )
        assert status == 400

    def test_04_加批号_指标不存在_404(self, web_server):
        base, _ = web_server
        status, body, _ = _http(
            "POST", f"{base}/api/test_items/99999/lots",
            data={"new_lot": "LOTX", "target": 100.0, "sd": 2.0},
        )
        assert status == 404


# ============================================================================
# §B.3 下载模板
# ============================================================================

class TestB3TemplateDownload:
    """spec §B.3：GET /api/template.xlsx → 6 必需列 + 1 行示例。"""

    def test_01_GET_template_xlsx(self, web_server):
        base, _ = web_server
        status, body, headers = _http("GET", f"{base}/api/template.xlsx")
        assert status == 200
        # xlsx 是 zip
        assert body[:4] == b"PK\x03\x04"
        # 解析
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(body))
        ws = wb.active
        # 读表头
        header = [c.value for c in ws[1]]
        for col in ("项目名", "检测指标", "水平", "批号", "日期", "测定值"):
            assert col in header, f"模板缺 {col} 列，实际 {header}"
        # 至少 1 行示例
        assert ws.max_row >= 2, f"模板应至少 1 行示例，实际 {ws.max_row} 行"


# ============================================================================
# §B.2 批量导入
# ============================================================================

class TestB2Import:
    """spec §B.2：POST /api/test_items/<tid>/import 上传 xlsx。"""

    def _make_xlsx(self, rows: list[dict]) -> bytes:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["项目名", "检测指标", "水平", "批号", "日期", "测定值", "序号"])
        for r in rows:
            ws.append([r["项目名"], r["检测指标"], r["水平"], r["批号"],
                       r["日期"], r["测定值"], r.get("序号", 0)])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_01_导入50行规范数据(self, web_server):
        base, tmp = web_server
        ids = _seed(base)
        rows = [
            {"项目名": "维生素 D", "检测指标": "D2", "水平": "L",
             "批号": "LOT1", "日期": f"2026-04-{(i % 28) + 1:02d}",
             "测定值": 100.0, "序号": i % 5}
            for i in range(50)
        ]
        xlsx = self._make_xlsx(rows)
        status, body, _ = _http(
            "POST", f"{base}/api/test_items/{ids['tid_l']}/import",
            raw=xlsx,
            headers={"X-Project-Name-Id": str(ids["pid"])},
        )
        assert status == 200, f"导入应 200，实际 {status}: {body}"
        result = json.loads(body)
        assert result["imported_count"] == 50, \
            f"应导入 50 行，实际 {result.get('imported_count')}"

    def test_02_导入_缺列_整批拒绝(self, web_server):
        base, _ = web_server
        ids = _seed(base)
        # 构造缺"测定值"列的 xlsx
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["项目名", "检测指标", "水平", "批号", "日期"])  # 缺"测定值"
        ws.append(["维生素 D", "D2", "L", "LOT1", "2026-04-01"])
        buf = io.BytesIO()
        wb.save(buf)
        xlsx = buf.getvalue()
        status, body, _ = _http(
            "POST", f"{base}/api/test_items/{ids['tid_l']}/import",
            raw=xlsx,
            headers={"X-Project-Name-Id": str(ids["pid"])},
        )
        assert status == 400
        err = json.loads(body).get("error", "")
        assert "测定值" in err or "缺少必需列" in err

    def test_03_导入_某行日期错_该行跳过(self, web_server):
        base, _ = web_server
        ids = _seed(base)
        rows = [
            {"项目名": "维生素 D", "检测指标": "D2", "水平": "L",
             "批号": "LOT1", "日期": "2026-04-01", "测定值": 100.0, "序号": 1},
            {"项目名": "维生素 D", "检测指标": "D2", "水平": "L",
             "批号": "LOT1", "日期": "不是日期", "测定值": 100.0, "序号": 2},  # 错
            {"项目名": "维生素 D", "检测指标": "D2", "水平": "L",
             "批号": "LOT1", "日期": "2026-04-03", "测定值": 100.0, "序号": 3},
        ]
        xlsx = self._make_xlsx(rows)
        status, body, _ = _http(
            "POST", f"{base}/api/test_items/{ids['tid_l']}/import",
            raw=xlsx,
            headers={"X-Project-Name-Id": str(ids["pid"])},
        )
        # 200 + skipped_rows 里有日期错
        result = json.loads(body)
        assert result["imported_count"] == 2, \
            f"应入 2 行，实际 {result.get('imported_count')}"
        assert any("日期" in s.get("reason", "") for s in result["skipped_rows"])

    def test_04_导入_项目不存在_整批拒绝(self, web_server):
        base, _ = web_server
        ids = _seed(base)
        rows = [
            {"项目名": "不存在的项目", "检测指标": "D2", "水平": "L",
             "批号": "LOT1", "日期": "2026-04-01", "测定值": 100.0, "序号": 1},
        ]
        xlsx = self._make_xlsx(rows)
        status, body, _ = _http(
            "POST", f"{base}/api/test_items/{ids['tid_l']}/import",
            raw=xlsx,
            headers={"X-Project-Name-Id": "99999"},
        )
        assert status == 400
        assert "项目" in json.loads(body).get("error", "")

    def test_05_导入_指标不存在_整批拒绝(self, web_server):
        base, _ = web_server
        ids = _seed(base)
        rows = [
            {"项目名": "维生素 D", "检测指标": "D3_不存在", "水平": "L",
             "批号": "LOT1", "日期": "2026-04-01", "测定值": 100.0, "序号": 1},
        ]
        xlsx = self._make_xlsx(rows)
        status, body, _ = _http(
            "POST", f"{base}/api/test_items/{ids['tid_l']}/import",
            raw=xlsx,
            headers={"X-Project-Name-Id": str(ids["pid"])},
        )
        assert status == 400
        assert "检测指标" in json.loads(body).get("error", "") or "D3" in json.loads(body).get("error", "")
