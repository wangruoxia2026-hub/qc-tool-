"""
测试：实验室、项目与检测指标（L/H 双水平）一体化建档
- HTML 包含一体化建档触发按钮与模态弹窗字段
- 包含实验室名称、项目信息、指标名称/批号、L 水平靶值/SD、H 水平靶值/SD
- 端到端验证：通过接口级联创建实验室、项目与 L/H 指标
"""
from __future__ import annotations

import json
import socket
import threading
import time
from http.server import ThreadingHTTPServer
from pathlib import Path
from urllib import request as urlreq

import pytest


def _free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


@pytest.fixture
def web_server(tmp_path: Path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    from qc_tool import web as webmod
    port = _free_port()
    handler = webmod.make_handler(db_path=tmp_path / "qc.db")
    server = ThreadingHTTPServer(("127.0.0.1", port), handler)
    server.db_path = tmp_path / "qc.db"
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base = f"http://127.0.0.1:{port}"
    time.sleep(0.05)
    try:
        yield base, tmp_path
    finally:
        server.shutdown()
        server.server_close()


def test_integrated_creation_html_elements(web_server):
    """验证 HTML 包含一体化建档入口与关键函数。"""
    base, _ = web_server
    r = urlreq.urlopen(f"{base}/", timeout=5)
    html = r.read().decode("utf-8")

    # 1. 按钮与弹窗函数
    assert 'openIntegratedModal' in html, "缺一体化建档入口函数 openIntegratedModal"
    assert '新建项目向导' in html, "缺 '新建项目向导' 按钮文案"
    assert 'saveIntegratedAll' in html, "缺一体化保存函数 saveIntegratedAll"
    # 2. 检查多指标与测算功能
    assert 'addIntegItemCard' in html, "缺多指标添加函数 addIntegItemCard"
    assert 'calculate_target_sd' in html, "缺 Excel 测算接口调用"
    assert 'onUploadMultiSheetCalcFile' in html, "缺多 Sheet Excel 测算函数 onUploadMultiSheetCalcFile"
    # 3. 检查检测人员一人一框
    assert 'addPersonRow' in html, "缺添加人员函数 addPersonRow"
    assert 'integ-pre-processor-list' in html, "缺前处理人员列表容器"
    assert 'integ-operator-list' in html, "缺仪器操作人员列表容器"


def test_integrated_creation_flow(web_server):
    """端到端验证级联创建：实验室 -> 项目 -> 指标L -> 指标H。"""
    base, tmp_path = web_server

    # 1. 保存实验室
    req = urlreq.Request(
        f"{base}/api/lab",
        data=json.dumps({"lab_name": "省人民医院检验科"}).encode("utf-8"),
        headers={"Content-Type": "application/json"},
    )
    resp = urlreq.urlopen(req, timeout=5)
    lab = json.loads(resp.read().decode("utf-8"))
    assert lab["lab_name"] == "省人民医院检验科"

    # 2. 创建项目
    req = urlreq.Request(
        f"{base}/api/projects",
        data=json.dumps({
            "name": "肝功能全套",
            "short": "LFT",
            "unit": "U/L",
            "manufacturer": "罗氏诊断",
        }).encode("utf-8"),
        headers={"Content-Type": "application/json"},
    )
    resp = urlreq.urlopen(req, timeout=5)
    proj = json.loads(resp.read().decode("utf-8"))
    pid = proj["id"]

    # 3. 创建 L 水平指标
    req_l = urlreq.Request(
        f"{base}/api/projects/{pid}/test_items",
        data=json.dumps({
            "name": "ALT",
            "level": "L",
            "lot": "LOT2026A",
            "target": 40.0,
            "sd": 2.0,
            "unit": "U/L",
            "manufacturer": "罗氏诊断",
        }).encode("utf-8"),
        headers={"Content-Type": "application/json"},
    )
    resp_l = urlreq.urlopen(req_l, timeout=5)
    item_l = json.loads(resp_l.read().decode("utf-8"))
    assert item_l["level"] == "L"
    assert item_l["target"] == 40.0

    # 4. 创建 H 水平指标
    req_h = urlreq.Request(
        f"{base}/api/projects/{pid}/test_items",
        data=json.dumps({
            "name": "ALT",
            "level": "H",
            "lot": "LOT2026A",
            "target": 120.0,
            "sd": 6.0,
            "unit": "U/L",
            "manufacturer": "罗氏诊断",
        }).encode("utf-8"),
        headers={"Content-Type": "application/json"},
    )
    resp_h = urlreq.urlopen(req_h, timeout=5)
    item_h = json.loads(resp_h.read().decode("utf-8"))
    assert item_h["level"] == "H"
    assert item_h["target"] == 120.0

    # 5. 查询该项目下是否包含这两个指标
    resp = urlreq.urlopen(f"{base}/api/projects/{pid}/test_items", timeout=5)
    items = json.loads(resp.read().decode("utf-8"))
    assert len(items) == 2
    levels = {it["level"] for it in items}
    assert levels == {"L", "H"}


def test_excel_target_sd_calculation_endpoint(web_server):
    """测试通过接口上传 Excel 文件自动测算靶值与 SD。"""
    import io
    from openpyxl import Workbook

    base, _ = web_server

    # 构造测试 Excel：包含水平列与测定值列
    # L: 10, 12, 14 -> Mean = 12.0, SD = 2.0
    # H: 100, 110, 120 -> Mean = 110.0, SD = 10.0
    wb = Workbook()
    ws = wb.active
    ws.append(["水平", "测定值"])
    ws.append(["L", 10.0])
    ws.append(["L", 12.0])
    ws.append(["L", 14.0])
    ws.append(["H", 100.0])
    ws.append(["H", 110.0])
    ws.append(["H", 120.0])
    bio = io.BytesIO()
    wb.save(bio)
    xlsx_bytes = bio.getvalue()

    req = urlreq.Request(
        f"{base}/api/calculate_target_sd",
        data=xlsx_bytes,
        headers={"Content-Type": "application/octet-stream"},
    )
    resp = urlreq.urlopen(req, timeout=5)
    res = json.loads(resp.read().decode("utf-8"))

    assert "L" in res and "H" in res
    assert res["L"]["n"] == 3
    assert res["L"]["mean"] == 12.0
    assert res["L"]["sd"] == 2.0
    assert res["H"]["n"] == 3
    assert res["H"]["mean"] == 110.0
    assert res["H"]["sd"] == 10.0


def test_excel_multi_sheet_target_sd_calculation(web_server):
    """测试多 Sheet Excel 测算：每个 Sheet 代表一个指标（如 ALT、AST），包含对应 L/H 测定数据。"""
    import io
    from openpyxl import Workbook

    base, _ = web_server

    wb = Workbook()
    # Sheet 1: ALT
    ws1 = wb.active
    ws1.title = "ALT"
    ws1.append(["水平", "测定值"])
    ws1.append(["L", 40.0])
    ws1.append(["L", 42.0])
    ws1.append(["L", 44.0])
    ws1.append(["H", 120.0])
    ws1.append(["H", 130.0])
    ws1.append(["H", 140.0])

    # Sheet 2: AST
    ws2 = wb.create_sheet(title="AST")
    ws2.append(["水平L", "水平H"])
    ws2.append([30.0, 90.0])
    ws2.append([32.0, 100.0])
    ws2.append([34.0, 110.0])

    bio = io.BytesIO()
    wb.save(bio)
    xlsx_bytes = bio.getvalue()

    req = urlreq.Request(
        f"{base}/api/calculate_target_sd",
        data=xlsx_bytes,
        headers={"Content-Type": "application/octet-stream"},
    )
    resp = urlreq.urlopen(req, timeout=5)
    res = json.loads(resp.read().decode("utf-8"))

    assert res["is_multi_sheet"] is True
    assert res["sheet_count"] == 2
    assert "ALT" in res["sheets"] and "AST" in res["sheets"]

    alt_data = res["sheets"]["ALT"]
    assert alt_data["L"]["mean"] == 42.0
    assert alt_data["L"]["sd"] == 2.0
    assert alt_data["H"]["mean"] == 130.0
    assert alt_data["H"]["sd"] == 10.0

    ast_data = res["sheets"]["AST"]
    assert ast_data["L"]["mean"] == 32.0
    assert ast_data["L"]["sd"] == 2.0
    assert ast_data["H"]["mean"] == 100.0
    assert ast_data["H"]["sd"] == 10.0


def test_user_template_format_calculation(web_server):
    """验证用户实际使用的模版结构：表头为 ['质控编号', 'levelL', 'levelH']，
    系统应能正确识别 levelL 作为 L 水平、levelH 作为 H 水平，跳过质控编号列。
    """
    import io
    from openpyxl import Workbook

    base, _ = web_server

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "维生素A"
    ws1.append(["质控编号", "levelL", "levelH"])
    ws1.append([2026101201, 10.0, 50.0])
    ws1.append([2026101202, 12.0, 52.0])
    ws1.append([2026101203, 14.0, 54.0])

    ws2 = wb.create_sheet(title="维生素E")
    ws2.append(["质控编号", "levelL", "levelH"])
    ws2.append([2026101201, 8.0, 40.0])
    ws2.append([2026101202, 10.0, 42.0])
    ws2.append([2026101203, 12.0, 44.0])

    bio = io.BytesIO()
    wb.save(bio)
    xlsx_bytes = bio.getvalue()

    # 通过真实 multipart/form-data POST 请求上传
    boundary = "----WebKitFormBoundaryUserTest123"
    body = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="file"; filename="上传模版.xlsx"\r\n'
        f"Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n"
    ).encode("utf-8") + xlsx_bytes + f"\r\n--{boundary}--\r\n".encode("utf-8")

    req = urlreq.Request(
        f"{base}/api/calculate_target_sd",
        data=body,
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
    )
    resp = urlreq.urlopen(req, timeout=5)
    res = json.loads(resp.read().decode("utf-8"))

    assert res["is_multi_sheet"] is True
    assert "维生素A" in res["sheets"] and "维生素E" in res["sheets"]

    va = res["sheets"]["维生素A"]
    # levelL 应正确被识别为 L 水平（而不是质控编号 2026101201）
    assert va["L"]["mean"] == 12.0
    assert va["L"]["sd"] == 2.0
    assert va["H"]["mean"] == 52.0
    assert va["H"]["sd"] == 2.0

    ve = res["sheets"]["维生素E"]
    assert ve["L"]["mean"] == 10.0
    assert ve["L"]["sd"] == 2.0
    assert ve["H"]["mean"] == 42.0
    assert ve["H"]["sd"] == 2.0

